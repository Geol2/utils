# egov_impact_report.py
# 사용: python egov_impact_report.py ./src TB_USER --start 2026-09-14 --people 2
#   pip install openpyxl
# 검색어: 테이블명 / mapper id / 컨트롤러 URL / JS 함수명(fn_xxx) 아무거나
import re, sys, time, argparse, pathlib, datetime as dt
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------- 공수 기준 (팀 기준으로 수정) ----------
# 항목당 기본 MD. 직접영향=1.0배, 간접영향=0.5배, 테스트는 별도 계산
BASE_MD = {
    'Table':      0.5,   # DDL/데이터 이관 검토
    'Mapper':     0.5,   # SQL 수정
    'DAO':        0.3,
    'Service':    0.5,
    'Controller': 0.5,
    'JSP':        0.7,   # 화면 수정 + 화면 확인
    'JS':         0.5,
}
TEST_RATIO   = 0.3     # 개발 공수 대비 테스트 공수 비율
BUFFER_RATIO = 0.15    # 리스크 버퍼
# ---------------------------------------------------

ap = argparse.ArgumentParser()
ap.add_argument('root'); ap.add_argument('query')
ap.add_argument('--start', default=dt.date.today().isoformat())
ap.add_argument('--people', type=float, default=1)
ap.add_argument('--out', default='영향도분석서.xlsx')
a = ap.parse_args()
root, q = pathlib.Path(a.root), a.query.lower()
read = lambda p: p.read_text(encoding='utf-8', errors='ignore')

# ---------- 진행 표시 ----------
try: sys.stdout.reconfigure(errors='replace')   # cp949 콘솔에서 ✔ 같은 기호 때문에 죽지 않게
except Exception: pass
_t0 = time.time(); _last = [0.0]; _cr = [False]
def _line(s):
    if _cr[0]: print(); _cr[0] = False
    print(s, flush=True)
def stage(s): _line(f'▶ {s}')
def done(s): _line(f'   ✔ {s}')
def tick(i, n, name):
    now = time.time()
    if i == n or now - _last[0] > 0.2:
        _last[0] = now; print('\r' + f'   {i}/{n}  {name}'[:110].ljust(110), end='', flush=True); _cr[0] = True

# ---------- 정규식 ----------
# 컨트롤러 메서드 매핑: @XxxMapping("url") [다른 어노테이션들] [제어자] 반환형 메서드명(
MAP_RE = re.compile(r'@(?:Request|Get|Post|Put|Delete)Mapping\(\s*(?:(?:value|path)\s*=\s*)?\{?\s*"([^"]+)"[^)]*\)\s*'
                    r'(?:@\w+(?:\([^)]*\))?\s*)*(?:(?:public|protected|private|static|final|synchronized)\s+)*[\w<>\[\], ?]+\s+(\w+)\s*\(')
CLASS_MAP_RE = re.compile(r'@RequestMapping\(\s*(?:(?:value|path)\s*=\s*)?\{?\s*"([^"]+)"')
# JS 함수 정의: function a(  /  a = function(  /  a: function(  /  a = (…) =>
FN_DEF_RE = re.compile(r'function\s+(\w+)\s*\(|\b(\w+)\s*=\s*(?:async\s+)?function\s*\(|\b(\w+)\s*:\s*(?:async\s+)?function\s*\(|\b(\w+)\s*=\s*(?:async\s*)?\([^()]*\)\s*=>')
FN_CALL_RE = re.compile(r'\b(\w+)\s*\(')
FN_SKIP = {'function', 'if', 'for', 'while', 'switch', 'catch', 'return', 'typeof'}
URL_LIT_RE = re.compile(r'["\']([^"\'\s<>]*/[^"\'\s<>]*)["\']')          # 슬래시가 든 문자열 리터럴 = URL 후보 ($.ajax url, form action, location.href, fetch 등 전부)
SCRIPT_SRC_RE = re.compile(r'<script[^>]+src\s*=\s*["\'][^>]*?([\w.-]+\.js)\b', re.I)   # src="<c:url value='/js/a.js'/>" 처럼 감싸도 파일명만 뽑음
JAVA_SUFFIX_RE = re.compile(r'(DAO|Dao|ServiceImpl|Service|Mapper)$')

def norm_url(u):
    u = re.sub(r'\$\{[^}]*\}|<%=[^%]*%>', '', u)          # ${ctx}, <%=ctx%> 제거
    return u.split('?')[0].split('#')[0]

def url_matcher(url):
    """컨트롤러 URL이 뷰 문자열 리터럴의 끝부분과 일치하는지. 컨텍스트 경로 prefix(/dWorks/user/list.do)와 {id} 같은 path variable 허용"""
    if not url.startswith('/'): url = '/' + url
    if '{' in url:
        pat = re.compile(re.sub(r'\\\{[^}]*\\\}', r'[^/]+', re.escape(url)) + r'$')
        return lambda lit: bool(pat.search(lit))
    return lambda lit: lit == url or lit.endswith(url)

def fn_spans(t):
    """JS/JSP 안의 함수 정의 → [(함수명, 시작, 끝)] (중괄호 짝으로 본문 범위 계산)"""
    stack, close = [], {}
    for i, ch in enumerate(t):
        if ch == '{': stack.append(i)
        elif ch == '}' and stack: close[stack.pop()] = i
    spans = []
    for m in FN_DEF_RE.finditer(t):
        fn = next(x for x in m.groups() if x)
        if len(fn) < 3 or fn in FN_SKIP: continue
        p = m.end()
        if m.group(4) is None:                       # 화살표 함수가 아니면 파라미터 목록 건너뜀
            p = t.find(')', p)
            if p < 0: continue
            p += 1
        while p < len(t) and t[p] in ' \t\r\n': p += 1
        if p < len(t) and t[p] == '{' and p in close: spans.append((fn, m.start(), close[p]))
        else: spans.append((fn, m.start(), t.find('\n', p) if t.find('\n', p) > 0 else len(t)))   # 한 줄짜리 화살표 함수
    return spans

def enclosing(spans, pos):
    best = None
    for fn, s, e in spans:
        if s <= pos <= e and (best is None or e - s < best[2] - best[1]): best = (fn, s, e)
    return best[0] if best else None

def kind_of(p):
    s = p.name
    if s.endswith('Controller.java'): return 'Controller'
    if 'DAO' in s or 'Dao' in s or 'Mapper.java' in s: return 'DAO'
    if s.endswith('.java'): return 'Service'
    if s.endswith('.jsp'): return 'JSP'
    if s.endswith('.js'): return 'JS'
    return 'Etc'

# ---------- 1. 인덱스 구축 / 2. 영향 항목 수집 ----------
if not root.is_dir(): sys.exit(f'폴더를 찾을 수 없습니다: {root}')

stage('파일 목록 수집')
java = list(root.rglob('*.java')); ctrls = [j for j in java if j.name.endswith('Controller.java')]
xmls = list(root.rglob('*.xml')); views = list(root.rglob('*.jsp')) + list(root.rglob('*.js'))
done(f'Java {len(java)}개 (Controller {len(ctrls)}) · XML {len(xmls)}개 · JSP/JS {len(views)}개')

url_to_method, id_to_tables, id_to_xml, lines = {}, {}, {}, {}
java_calls_id = defaultdict(set)      # java 파일 -> mapper ids
view_calls_url = defaultdict(set)     # url -> {(파일, 함수명|None)}
fn_urls = defaultdict(set)            # (파일, 함수명) -> urls
fn_defs = defaultdict(set)            # 함수명 -> 정의한 파일
fn_callers = defaultdict(set)         # 함수명 -> {(호출 파일, 호출한 함수명|None)} (정의 파일 제외)
js_included_by = defaultdict(set)     # js 파일명 -> <script src> 로 include 한 jsp

stage('Controller URL 매핑')
for i, j in enumerate(ctrls, 1):
    tick(i, len(ctrls), j.name); t = read(j)
    m = re.search(r'\bclass\s+\w+', t)                      # 클래스 선언 앞에 있는 @RequestMapping 만 클래스 레벨 prefix
    prefix = CLASS_MAP_RE.search(t[:m.start()]) if m else None
    prefix = prefix.group(1) if prefix else ''
    for url, name in MAP_RE.findall(t): url_to_method[(prefix + url).replace('//', '/')] = (j, name)
done(f'URL {len(url_to_method)}개')

stage('Mapper XML 파싱')
for i, x in enumerate(xmls, 1):
    tick(i, len(xmls), x.name); t = read(x)
    for _, sid, body in re.findall(r'<(select|insert|update|delete)\s+id="([^"]+)"(.*?)</\1>', t, re.S):
        id_to_tables[sid] = {m.upper() for m in re.findall(r'\b(?:FROM|JOIN|INTO|UPDATE)\s+([A-Za-z_][\w.]*)', body, re.I)}
        id_to_xml[sid] = x
done(f'statement {len(id_to_tables)}개 · 테이블 {len(set().union(*id_to_tables.values()))}개')

stage('Java → Mapper 호출 추적')
sid_re = re.compile(r'["\.](' + '|'.join(sorted(map(re.escape, id_to_tables), key=len, reverse=True)) + r')\s*["(]') if id_to_tables else None
for i, j in enumerate(java, 1):
    tick(i, len(java), j.name); t = read(j); lines[j] = t.count('\n')
    if sid_re: java_calls_id[j].update(sid_re.findall(t))
done(f'mapper 를 호출하는 Java {sum(1 for v in java_calls_id.values() if v)}개')

stage('JSP/JS 스캔 (URL 호출 · 함수 정의 · script include)')
url_pats = []
for u in url_to_method:
    seg = u.rsplit('/', 1)[-1]
    url_pats.append((u, '' if '{' in seg else seg, url_matcher(u)))   # seg: 빠른 사전 필터용 (마지막 경로 조각)
view_text, view_spans = {}, {}
for i, f in enumerate(views, 1):
    tick(i, len(views), f.name); t = read(f); view_text[f] = t; lines[f] = t.count('\n')
    spans = fn_spans(t); view_spans[f] = spans
    for fn, _, _ in spans: fn_defs[fn].add(f)
    lits = [(norm_url(m.group(1)), m.start(1)) for m in URL_LIT_RE.finditer(t)]
    for u, seg, match in url_pats:
        if seg not in t: continue
        for lit, pos in lits:
            if match(lit):
                fn = enclosing(spans, pos)
                view_calls_url[u].add((f, fn)); fn_urls[(f, fn)].add(u)
    if f.suffix.lower() == '.jsp':
        for src in SCRIPT_SRC_RE.findall(t): js_included_by[src].add(f)
for f, t in view_text.items():
    for m in FN_CALL_RE.finditer(t):
        fn = m.group(1)
        if fn in fn_defs and f not in fn_defs[fn]: fn_callers[fn].add((f, enclosing(view_spans[f], m.start())))
done(f'URL 호출 {sum(len(v) for v in view_calls_url.values())}건 · 함수 정의 {len(fn_defs)}개 · script include {sum(len(v) for v in js_included_by.values())}건')

stage('영향 항목 수집')
items, nodes, edges = {}, {}, set()
def add(kind, path, name, impact, via=''):
    k = (kind, str(path), name)
    if k not in items or impact == '직접': items[k] = {'impact': impact, 'via': via}
def node(kind, key, impact, label=None, title=''):
    nid = f'{kind}|{key}'; n = nodes.get(nid)
    if n is None: nodes[nid] = {'id': nid, 'kind': kind, 'label': label or key, 'impact': impact, 'title': title}
    elif impact == '직접': n['impact'] = '직접'
    return nid
def link(a, b, rel=''): edges.add((a, b, rel))
def vnode(f, fn, impact):
    return node(kind_of(f), f'{f.name}::{fn}' if fn else f.name, impact, f'{f.name} › {fn}()' if fn else f.name, str(f))
def cnode(cj, m, url, impact):
    return node('Controller', f'{cj.stem}.{m} ({url})', impact, f'{cj.stem}.{m}', f'{url}\n{cj}')

def hit_view(f, fn, impact, url, cn):
    """뷰(JSP/JS)가 컨트롤러 URL 을 부름 → 항목 추가 + 그 함수를 호출하는 곳까지 한 단계 더"""
    add(kind_of(f), f, f.name, impact, f'{url} ← {fn}()' if fn else url)
    vn = vnode(f, fn, impact); link(vn, cn)
    if fn:
        for cf, cfn in fn_callers.get(fn, []):
            add(kind_of(cf), cf, cf.name, '간접', f'{fn}() 호출'); link(vnode(cf, cfn, '간접'), vn)
def hit_ctrl(cj, m, url, impact, via=''):
    add('Controller', cj, f'{cj.stem}.{m} ({url})', impact, via)
    cn = cnode(cj, m, url, impact)
    for v, fn in view_calls_url.get(url, []): hit_view(v, fn, impact, url, cn)
    return cn
traced = set()
def trace_down(cn, cj):
    """컨트롤러 → (클래스명 관례로) Service/DAO → 그 클래스가 쓰는 mapper → 테이블. 클래스 단위 근사라 전부 간접"""
    if cn in traced: return
    traced.add(cn); cls = re.sub(r'Controller$', '', cj.stem)
    g = {'Service': [], 'DAO': []}
    for j, ids in java_calls_id.items():
        if ids and JAVA_SUFFIX_RE.sub('', j.stem) == cls:
            add(kind_of(j), j, j.stem, '간접', cj.stem)
            jn = node(kind_of(j), j.stem, '간접', title=str(j)); g['DAO' if kind_of(j) == 'DAO' else 'Service'].append(jn)
            for sid in ids:
                add('Mapper', id_to_xml[sid], sid, '간접', j.stem)
                mn = node('Mapper', sid, '간접', title=str(id_to_xml[sid])); link(jn, mn)
                for t in id_to_tables[sid]: add('Table', '-', t, '간접', sid); link(mn, node('Table', t, '간접'))
    for s in g['Service']:
        for d in g['DAO']: link(s, d)
    for x in (g['Service'] or g['DAO']): link(cn, x)

# 테이블 / mapper id 기준: 아래에서 위로
for sid, tables in id_to_tables.items():
    hit_tbl = [t for t in tables if q in t.lower()]
    if hit_tbl or q in sid.lower():
        mn = node('Mapper', sid, '직접', title=str(id_to_xml[sid]))
        for t in hit_tbl: add('Table', '-', t, '직접'); link(mn, node('Table', t, '직접'))
        add('Mapper', id_to_xml[sid], sid, '직접', ','.join(hit_tbl))
        by_cls = defaultdict(lambda: {'Service': [], 'DAO': [], 'stems': []})
        for j, ids in java_calls_id.items():
            if sid in ids:
                add(kind_of(j), j, j.stem, '직접', sid)
                jn = node(kind_of(j), j.stem, '직접', title=str(j)); link(jn, mn)
                cls = JAVA_SUFFIX_RE.sub('', j.stem)
                if cls: g = by_cls[cls]; g['DAO' if kind_of(j) == 'DAO' else 'Service'].append(jn); g['stems'].append(j.stem)
        for cls, g in by_cls.items():
            for s in g['Service']:
                for d in g['DAO']: link(s, d)
            for url, (cj, m) in url_to_method.items():
                if cls in cj.stem:
                    cn = hit_ctrl(cj, m, url, '간접', ', '.join(g['stems']))
                    for x in (g['Service'] or g['DAO']): link(cn, x)
done(f'테이블/Mapper 기준 {len(items)}건'); n = len(items)
# URL 기준: 컨트롤러에서 양쪽으로
for url, (cj, m) in url_to_method.items():
    if q in url.lower() or q in m.lower(): trace_down(hit_ctrl(cj, m, url, '직접'), cj)
done(f'URL 기준 +{len(items) - n}건'); n = len(items)
# JS 함수 기준 (식별자 형태 검색어일 때만): 정의한 파일 직접, 호출한 파일 간접, 그 함수가 부르는 컨트롤러부터 아래로
if re.fullmatch(r'\w+', a.query):
    for fn, files in fn_defs.items():
        if q in fn.lower():
            for v in files:
                add(kind_of(v), v, f'{v.name}::{fn}', '직접', f'{fn} 정의'); vn = vnode(v, fn, '직접')
                for url in fn_urls.get((v, fn), []):
                    cj, m = url_to_method[url]
                    add('Controller', cj, f'{cj.stem}.{m} ({url})', '간접', f'{fn}()'); cn = cnode(cj, m, url, '간접')
                    link(vn, cn); trace_down(cn, cj)
                for cf, cfn in fn_callers.get(fn, []):
                    add(kind_of(cf), cf, cf.name, '간접', f'{fn}() 호출'); link(vnode(cf, cfn, '간접'), vn)
done(f'JS 함수 기준 +{len(items) - n}건'); n = len(items)
# 영향받은 JS 를 <script src> 로 물고 있는 JSP
for jn in [x for x in nodes.values() if x['kind'] == 'JS']:
    for jsp in js_included_by.get(jn['id'].split('|', 1)[1].split('::')[0], []):
        add('JSP', jsp, jsp.name, '간접', f'{jn["label"].split(" ›")[0]} include'); link(vnode(jsp, None, '간접'), jn['id'], 'include')
done(f'JS 를 include 하는 JSP +{len(items) - n}건')

# 관계 정리: 같은 쌍에 호출/include 가 둘 다 있으면 호출만, 중간 노드를 거쳐 갈 수 있는 지름길 간선은 제거
pairs = {(x, y) for x, y, _ in edges}
edges = {(x, y, r) for x, y, r in edges if not (r and (x, y, '') in edges)}
out = defaultdict(set)
for x, y in pairs: out[x].add(y)
edges = {(x, y, r) for x, y, r in edges if r or not any(y in out[m] for m in out[x] if m != y)}

stage('공수·일정 산정')

# ---------- 3. 공수/일정 ----------
def md_of(kind, path, impact):
    base = BASE_MD.get(kind, 0.3)
    p = pathlib.Path(path)
    ln = lines.get(p, 0)
    size = 1.0 if ln < 300 else 1.3 if ln < 1000 else 1.6   # 파일 크기 가중
    return round(base * size * (1.0 if impact == '직접' else 0.5), 2)

rows = []
for (kind, path, name), v in sorted(items.items(), key=lambda kv: (list(BASE_MD).index(kv[0][0]) if kv[0][0] in BASE_MD else 9, kv[0][2])):
    rows.append([kind, name, path if path != '-' else '', v['impact'], v['via'], md_of(kind, path, v['impact'])])

dev_md  = round(sum(r[5] for r in rows), 2)
test_md = round(dev_md * TEST_RATIO, 2)
buf_md  = round((dev_md + test_md) * BUFFER_RATIO, 2)
total   = round(dev_md + test_md + buf_md, 2)

def add_workdays(d, n):
    while n > 0:
        d += dt.timedelta(days=1)
        if d.weekday() < 5: n -= 1
    return d
start = dt.date.fromisoformat(a.start)
if start.weekday() >= 5: start = add_workdays(start, 1)
days = -(-total // a.people)  # ceil
end = add_workdays(start, max(int(days) - 1, 0))

# 항목별 일정: 순차 배분 (투입인원 반영)
cur = start; acc = 0.0
for r in rows:
    s = cur
    acc += r[5]
    e = add_workdays(start, max(int(-(-acc // a.people)) - 1, 0))
    r += [s.isoformat(), e.isoformat()]
    cur = e

# ---------- 4. 엑셀 ----------
wb = Workbook()
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='305496')
thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)

def sheet(ws, header, data, widths):
    ws.append(header)
    for c in ws[1]: c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center'); c.border = B
    for row in data:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.border = B
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'

ws = wb.active; ws.title = '요약'
sheet(ws, ['항목', '값'], [
    ['변경 대상', a.query], ['분석 일자', dt.date.today().isoformat()],
    ['영향 항목 수', len(rows)],
    ['직접 영향', sum(1 for r in rows if r[3] == '직접')], ['간접 영향', sum(1 for r in rows if r[3] == '간접')],
    ['개발 공수 (MD)', dev_md], [f'테스트 공수 (MD, {int(TEST_RATIO*100)}%)', test_md],
    [f'버퍼 (MD, {int(BUFFER_RATIO*100)}%)', buf_md], ['총 공수 (MD)', total],
    ['투입 인원', a.people], ['예상 기간 (영업일)', int(days)],
    ['시작일', start.isoformat()], ['종료일', end.isoformat()],
], [28, 20])

ws2 = wb.create_sheet('영향목록')
sheet(ws2, ['No', '구분', '항목', '파일', '영향', '경유', '공수(MD)', '시작일', '종료일', '담당자', '조치내용', '비고'],
      [[i + 1] + r + ['', '', ''] for i, r in enumerate(rows)],
      [5, 11, 40, 55, 7, 22, 9, 11, 11, 10, 30, 20])
ws2.auto_filter.ref = ws2.dimensions

lab = nodes
ws4 = wb.create_sheet('연결관계')
sheet(ws4, ['From 구분', 'From', 'To 구분', 'To', '관계'],
      [[lab[x]['kind'], lab[x]['label'], lab[y]['kind'], lab[y]['label'], 'include' if r else '참조' if lab[y]['kind'] == 'Table' else '호출'] for x, y, r in sorted(edges)], [11, 40, 11, 40, 9])
ws4.auto_filter.ref = ws4.dimensions

ws3 = wb.create_sheet('공수기준')
sheet(ws3, ['구분', '기본 MD/항목', '비고'],
      [[k, v, ''] for k, v in BASE_MD.items()] +
      [['파일크기 가중', '1.0 / 1.3 / 1.6', '<300 / <1000 / 1000+ 라인'],
       ['간접영향', 'x0.5', ''], ['테스트', f'x{TEST_RATIO}', '개발공수 대비'], ['버퍼', f'x{BUFFER_RATIO}', '개발+테스트 대비']],
      [16, 16, 30])

wb.save(a.out)
done(f'{a.out} 저장 — 영향 {len(rows)}건, 관계 {len(edges)}개, 총 {total} MD, {start}~{end} ({int(days)} 영업일 / {a.people}명), {time.time() - _t0:.1f}초')
