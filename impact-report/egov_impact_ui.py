# egov_impact_ui.py — 스프링 영향도 분석 로컬 UI
# 실행: python egov_impact_ui.py   (브라우저가 자동으로 열림, http://127.0.0.1:8765)
# 의존성: openpyxl (엑셀 저장에만 사용). 나머지는 표준 라이브러리.
import re, io, sys, json, time, pathlib, datetime as dt, webbrowser, threading
from collections import defaultdict
from http.server import HTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse

PORT = 8765
HIST_FILE = pathlib.Path(__file__).with_name('impact_history.json')
HIST_MAX = 30

# ---------- 공수 기준 (팀 기준으로 수정) ----------
BASE_MD = {'Table': 0.5, 'Mapper': 0.5, 'DAO': 0.3, 'Service': 0.5, 'Controller': 0.5, 'JSP': 0.7, 'JS': 0.5}
TEST_RATIO, BUFFER_RATIO = 0.3, 0.15
# ---------------------------------------------------

read = lambda p: p.read_text(encoding='utf-8', errors='ignore')

def add_workdays(d, n):
    while n > 0:
        d += dt.timedelta(days=1)
        if d.weekday() < 5: n -= 1
    return d

def load_hist():
    try: return json.loads(HIST_FILE.read_text(encoding='utf-8'))
    except Exception: return []

def save_hist(entry):
    h = [x for x in load_hist() if not (x['root'] == entry['root'] and x['query'] == entry['query'])]
    h.insert(0, entry)
    HIST_FILE.write_text(json.dumps(h[:HIST_MAX], ensure_ascii=False, indent=1), encoding='utf-8')
    return h[:HIST_MAX]

# ---------- 정규식 ----------
# 컨트롤러 메서드 매핑: @XxxMapping("url") [다른 어노테이션들] [제어자] 반환형 메서드명(
MAP_RE = re.compile(r'@(?:Request|Get|Post|Put|Delete)Mapping\(\s*(?:(?:value|path)\s*=\s*)?\{?\s*"([^"]+)"[^)]*\)\s*'
                    r'(?:@\w+(?:\([^)]*\))?\s*)*(?:(?:public|protected|private|static|final|synchronized)\s+)*[\w<>\[\], ?]+\s+(\w+)\s*\(')
CLASS_MAP_RE = re.compile(r'@RequestMapping\(\s*(?:(?:value|path)\s*=\s*)?\{?\s*"([^"]+)"')
# JS 함수 정의: function a(  /  a = function(  /  a: function(  /  a = (…) =>
FN_DEF_RE = re.compile(r'function\s+(\w+)\s*\(|\b(\w+)\s*=\s*(?:async\s+)?function\s*\(|\b(\w+)\s*:\s*(?:async\s+)?function\s*\(|\b(\w+)\s*=\s*(?:async\s*)?\([^()]*\)\s*=>')
FN_CALL_RE = re.compile(r'\b(\w+)\s*\(')
FN_SKIP = {'function', 'if', 'for', 'while', 'switch', 'catch', 'return', 'typeof'}
URL_LIT_RE = re.compile(r'["\']([^"\'\s<>]*/[^"\'\s<>]*)["\']')          # 슬래시가 든 문자열 리터럴 = URL 후보 ($.ajax url, form action, location.href, fetch 등 전부)
SCRIPT_SRC_RE = re.compile(r'<script[^>]+src\s*=\s*["\'][^>]*?([\w.-]+\.js)\b', re.I)   # src="<c:url value='/js/a.js'/>" 처럼 감싸도 파일명만 뽑음
JAVA_SUFFIX_RE = re.compile(r'(DAO|Dao|ServiceImpl|Service|Mapper)$')

def norm_url(u):
    u = re.sub(r'\$\{[^}]*\}|<%=[^%]*%>', '', u)          # ${ctx}, <%=ctx%> 제거
    return u.split('?')[0].split('#')[0]

def url_matcher(url):
    """컨트롤러 URL이 뷰 문자열 리터럴의 끝부분과 일치하는지. 컨텍스트 경로 prefix(/dWorks/user/list.do)와 {id} 같은 path variable 허용"""
    if not url.startswith('/'): url = '/' + url
    if '{' in url:
        pat = re.compile(re.sub(r'\\\{[^}]*\\\}', r'[^/]+', re.escape(url)) + r'$')
        return lambda lit: bool(pat.search(lit))
    return lambda lit: lit == url or lit.endswith(url)

def fn_spans(t):
    """JS/JSP 안의 함수 정의 → [(함수명, 시작, 끝)] (중괄호 짝으로 본문 범위 계산)"""
    stack, close = [], {}
    for i, ch in enumerate(t):
        if ch == '{': stack.append(i)
        elif ch == '}' and stack: close[stack.pop()] = i
    spans = []
    for m in FN_DEF_RE.finditer(t):
        fn = next(x for x in m.groups() if x)
        if len(fn) < 3 or fn in FN_SKIP: continue
        p = m.end()
        if m.group(4) is None:                       # 화살표 함수가 아니면 파라미터 목록 건너뜀
            p = t.find(')', p)
            if p < 0: continue
            p += 1
        while p < len(t) and t[p] in ' \t\r\n': p += 1
        if p < len(t) and t[p] == '{' and p in close: spans.append((fn, m.start(), close[p]))
        else: spans.append((fn, m.start(), t.find('\n', p) if t.find('\n', p) > 0 else len(t)))   # 한 줄짜리 화살표 함수
    return spans

def enclosing(spans, pos):
    best = None
    for fn, s, e in spans:
        if s <= pos <= e and (best is None or e - s < best[2] - best[1]): best = (fn, s, e)
    return best[0] if best else None

def kind_of(p):
    s = p.name
    if s.endswith('Controller.java'): return 'Controller'
    if 'DAO' in s or 'Dao' in s or 'Mapper.java' in s: return 'DAO'
    if s.endswith('.java'): return 'Service'
    if s.endswith('.jsp'): return 'JSP'
    if s.endswith('.js'): return 'JS'
    return 'Etc'

def analyze(root, query, start, people, log=lambda s, replace=False: None):
    root = pathlib.Path(root); q = query.lower()
    if not root.is_dir(): raise ValueError(f'폴더를 찾을 수 없습니다: {root}')
    t0 = time.time(); last = [0.0]
    def stage(s): log(f'▶ {s}')
    def done(s): log(f'   ✔ {s}')
    def tick(i, n, name):
        now = time.time()
        if i == n or now - last[0] > 0.2: last[0] = now; log(f'   {i}/{n}  {name}', replace=True)

    stage('파일 목록 수집')
    java = list(root.rglob('*.java')); ctrls = [j for j in java if j.name.endswith('Controller.java')]
    xmls = list(root.rglob('*.xml')); views = list(root.rglob('*.jsp')) + list(root.rglob('*.js'))
    done(f'Java {len(java)}개 (Controller {len(ctrls)}) · XML {len(xmls)}개 · JSP/JS {len(views)}개')

    url_to_method, id_to_tables, id_to_xml, lines = {}, {}, {}, {}
    java_calls_id = defaultdict(set)      # java 파일 -> mapper ids
    view_calls_url = defaultdict(set)     # url -> {(파일, 함수명|None)}
    fn_urls = defaultdict(set)            # (파일, 함수명) -> urls
    fn_defs = defaultdict(set)            # 함수명 -> 정의한 파일
    fn_callers = defaultdict(set)         # 함수명 -> {(호출 파일, 호출한 함수명|None)} (정의 파일 제외)
    js_included_by = defaultdict(set)     # js 파일명 -> <script src> 로 include 한 jsp

    stage('Controller URL 매핑')
    for i, j in enumerate(ctrls, 1):
        tick(i, len(ctrls), j.name); t = read(j)
        m = re.search(r'\bclass\s+\w+', t)                      # 클래스 선언 앞에 있는 @RequestMapping 만 클래스 레벨 prefix
        prefix = CLASS_MAP_RE.search(t[:m.start()]) if m else None
        prefix = prefix.group(1) if prefix else ''
        for url, name in MAP_RE.findall(t): url_to_method[(prefix + url).replace('//', '/')] = (j, name)
    done(f'URL {len(url_to_method)}개')

    stage('Mapper XML 파싱')
    for i, x in enumerate(xmls, 1):
        tick(i, len(xmls), x.name); t = read(x)
        for _, sid, body in re.findall(r'<(select|insert|update|delete)\s+id="([^"]+)"(.*?)</\1>', t, re.S):
            id_to_tables[sid] = {m.upper() for m in re.findall(r'\b(?:FROM|JOIN|INTO|UPDATE)\s+([A-Za-z_][\w.]*)', body, re.I)}
            id_to_xml[sid] = x
    done(f'statement {len(id_to_tables)}개 · 테이블 {len(set().union(*id_to_tables.values()))}개')

    stage('Java → Mapper 호출 추적')
    sid_re = re.compile(r'["\.](' + '|'.join(sorted(map(re.escape, id_to_tables), key=len, reverse=True)) + r')\s*["(]') if id_to_tables else None
    for i, j in enumerate(java, 1):
        tick(i, len(java), j.name); t = read(j); lines[j] = t.count('\n')
        if sid_re: java_calls_id[j].update(sid_re.findall(t))
    done(f'mapper 를 호출하는 Java {sum(1 for v in java_calls_id.values() if v)}개')

    stage('JSP/JS 스캔 (URL 호출 · 함수 정의 · script include)')
    url_pats = []
    for u in url_to_method:
        seg = u.rsplit('/', 1)[-1]
        url_pats.append((u, '' if '{' in seg else seg, url_matcher(u)))   # seg: 빠른 사전 필터용 (마지막 경로 조각)
    view_text, view_spans = {}, {}
    for i, f in enumerate(views, 1):
        tick(i, len(views), f.name); t = read(f); view_text[f] = t; lines[f] = t.count('\n')
        spans = fn_spans(t); view_spans[f] = spans
        for fn, _, _ in spans: fn_defs[fn].add(f)
        lits = [(norm_url(m.group(1)), m.start(1)) for m in URL_LIT_RE.finditer(t)]
        for u, seg, match in url_pats:
            if seg not in t: continue
            for lit, pos in lits:
                if match(lit):
                    fn = enclosing(spans, pos)
                    view_calls_url[u].add((f, fn)); fn_urls[(f, fn)].add(u)
        if f.suffix.lower() == '.jsp':
            for src in SCRIPT_SRC_RE.findall(t): js_included_by[src].add(f)
    for f, t in view_text.items():
        for m in FN_CALL_RE.finditer(t):
            fn = m.group(1)
            if fn in fn_defs and f not in fn_defs[fn]: fn_callers[fn].add((f, enclosing(view_spans[f], m.start())))
    done(f'URL 호출 {sum(len(v) for v in view_calls_url.values())}건 · 함수 정의 {len(fn_defs)}개 · script include {sum(len(v) for v in js_included_by.values())}건')

    stage('영향 항목 수집')
    items, nodes, edges = {}, {}, set()
    def add(kind, path, name, impact, via=''):
        k = (kind, str(path), name)
        if k not in items or impact == '직접': items[k] = {'impact': impact, 'via': via}
    def node(kind, key, impact, label=None, title=''):
        nid = f'{kind}|{key}'; n = nodes.get(nid)
        if n is None: nodes[nid] = {'id': nid, 'kind': kind, 'label': label or key, 'impact': impact, 'title': title}
        elif impact == '직접': n['impact'] = '직접'
        return nid
    def link(a, b, rel=''): edges.add((a, b, rel))
    def vnode(f, fn, impact):
        return node(kind_of(f), f'{f.name}::{fn}' if fn else f.name, impact, f'{f.name} › {fn}()' if fn else f.name, str(f))
    def cnode(cj, m, url, impact):
        return node('Controller', f'{cj.stem}.{m} ({url})', impact, f'{cj.stem}.{m}', f'{url}\n{cj}')

    def hit_view(f, fn, impact, url, cn):
        """뷰(JSP/JS)가 컨트롤러 URL 을 부름 → 항목 추가 + 그 함수를 호출하는 곳까지 한 단계 더"""
        add(kind_of(f), f, f.name, impact, f'{url} ← {fn}()' if fn else url)
        vn = vnode(f, fn, impact); link(vn, cn)
        if fn:
            for cf, cfn in fn_callers.get(fn, []):
                add(kind_of(cf), cf, cf.name, '간접', f'{fn}() 호출'); link(vnode(cf, cfn, '간접'), vn)
    def hit_ctrl(cj, m, url, impact, via=''):
        add('Controller', cj, f'{cj.stem}.{m} ({url})', impact, via)
        cn = cnode(cj, m, url, impact)
        for v, fn in view_calls_url.get(url, []): hit_view(v, fn, impact, url, cn)
        return cn
    traced = set()
    def trace_down(cn, cj):
        """컨트롤러 → (클래스명 관례로) Service/DAO → 그 클래스가 쓰는 mapper → 테이블. 클래스 단위 근사라 전부 간접"""
        if cn in traced: return
        traced.add(cn); cls = re.sub(r'Controller$', '', cj.stem)
        g = {'Service': [], 'DAO': []}
        for j, ids in java_calls_id.items():
            if ids and JAVA_SUFFIX_RE.sub('', j.stem) == cls:
                add(kind_of(j), j, j.stem, '간접', cj.stem)
                jn = node(kind_of(j), j.stem, '간접', title=str(j)); g['DAO' if kind_of(j) == 'DAO' else 'Service'].append(jn)
                for sid in ids:
                    add('Mapper', id_to_xml[sid], sid, '간접', j.stem)
                    mn = node('Mapper', sid, '간접', title=str(id_to_xml[sid])); link(jn, mn)
                    for t in id_to_tables[sid]: add('Table', '-', t, '간접', sid); link(mn, node('Table', t, '간접'))
        for s in g['Service']:
            for d in g['DAO']: link(s, d)
        for x in (g['Service'] or g['DAO']): link(cn, x)

    # 테이블 / mapper id 기준: 아래에서 위로
    for sid, tables in id_to_tables.items():
        hit_tbl = [t for t in tables if q in t.lower()]
        if hit_tbl or q in sid.lower():
            mn = node('Mapper', sid, '직접', title=str(id_to_xml[sid]))
            for t in hit_tbl: add('Table', '-', t, '직접'); link(mn, node('Table', t, '직접'))
            add('Mapper', id_to_xml[sid], sid, '직접', ','.join(hit_tbl))
            by_cls = defaultdict(lambda: {'Service': [], 'DAO': [], 'stems': []})
            for j, ids in java_calls_id.items():
                if sid in ids:
                    add(kind_of(j), j, j.stem, '직접', sid)
                    jn = node(kind_of(j), j.stem, '직접', title=str(j)); link(jn, mn)
                    cls = JAVA_SUFFIX_RE.sub('', j.stem)
                    if cls: g = by_cls[cls]; g['DAO' if kind_of(j) == 'DAO' else 'Service'].append(jn); g['stems'].append(j.stem)
            for cls, g in by_cls.items():
                for s in g['Service']:
                    for d in g['DAO']: link(s, d)
                for url, (cj, m) in url_to_method.items():
                    if cls in cj.stem:
                        cn = hit_ctrl(cj, m, url, '간접', ', '.join(g['stems']))
                        for x in (g['Service'] or g['DAO']): link(cn, x)
    done(f'테이블/Mapper 기준 {len(items)}건'); n = len(items)
    # URL 기준: 컨트롤러에서 양쪽으로
    for url, (cj, m) in url_to_method.items():
        if q in url.lower() or q in m.lower(): trace_down(hit_ctrl(cj, m, url, '직접'), cj)
    done(f'URL 기준 +{len(items) - n}건'); n = len(items)
    # JS 함수 기준 (식별자 형태 검색어일 때만): 정의한 파일 직접, 호출한 파일 간접, 그 함수가 부르는 컨트롤러부터 아래로
    if re.fullmatch(r'\w+', query):
        for fn, files in fn_defs.items():
            if q in fn.lower():
                for v in files:
                    add(kind_of(v), v, f'{v.name}::{fn}', '직접', f'{fn} 정의'); vn = vnode(v, fn, '직접')
                    for url in fn_urls.get((v, fn), []):
                        cj, m = url_to_method[url]
                        add('Controller', cj, f'{cj.stem}.{m} ({url})', '간접', f'{fn}()'); cn = cnode(cj, m, url, '간접')
                        link(vn, cn); trace_down(cn, cj)
                    for cf, cfn in fn_callers.get(fn, []):
                        add(kind_of(cf), cf, cf.name, '간접', f'{fn}() 호출'); link(vnode(cf, cfn, '간접'), vn)
    done(f'JS 함수 기준 +{len(items) - n}건'); n = len(items)
    # 영향받은 JS 를 <script src> 로 물고 있는 JSP
    for jn in [x for x in nodes.values() if x['kind'] == 'JS']:
        for jsp in js_included_by.get(jn['id'].split('|', 1)[1].split('::')[0], []):
            add('JSP', jsp, jsp.name, '간접', f'{jn["label"].split(" ›")[0]} include'); link(vnode(jsp, None, '간접'), jn['id'], 'include')
    done(f'JS 를 include 하는 JSP +{len(items) - n}건')
    # 관계 정리: 같은 쌍에 호출/include 가 둘 다 있으면 호출만, 중간 노드를 거쳐 갈 수 있는 지름길 간선(C→DAO 가 있는데 C→Service→DAO 도 있음)은 제거
    pairs = {(a, b) for a, b, _ in edges}
    edges = {(a, b, r) for a, b, r in edges if not (r and (a, b, '') in edges)}
    out = defaultdict(set)
    for a, b in pairs: out[a].add(b)
    edges = {(a, b, r) for a, b, r in edges if r or not any(b in out[m] for m in out[a] if m != b)}

    stage('공수·일정 산정')
    def md_of(kind, path, impact):
        ln = lines.get(pathlib.Path(path), 0)
        size = 1.0 if ln < 300 else 1.3 if ln < 1000 else 1.6
        return round(BASE_MD.get(kind, 0.3) * size * (1.0 if impact == '직접' else 0.5), 2)

    order = list(BASE_MD)
    rows = []
    for (kind, path, name), v in sorted(items.items(), key=lambda kv: (order.index(kv[0][0]) if kv[0][0] in order else 9, kv[0][2])):
        rows.append({'kind': kind, 'name': name, 'path': '' if path == '-' else path,
                     'impact': v['impact'], 'via': v['via'], 'md': md_of(kind, path, v['impact'])})

    dev = round(sum(r['md'] for r in rows), 2); test = round(dev * TEST_RATIO, 2)
    buf = round((dev + test) * BUFFER_RATIO, 2); total = round(dev + test + buf, 2)
    s = dt.date.fromisoformat(start)
    if s.weekday() >= 5: s = add_workdays(s, 1)
    days = int(-(-total // people)) if total else 0
    end = add_workdays(s, max(days - 1, 0))
    acc = 0.0; cur = s
    for r in rows:
        acc += r['md']; e = add_workdays(s, max(int(-(-acc // people)) - 1, 0))
        r['start'], r['end'] = cur.isoformat(), e.isoformat(); cur = e
    done(f'완료 — 영향 {len(rows)}건, 관계 {len(edges)}개, 총 {total} MD, {time.time() - t0:.1f}초')
    return {'query': query, 'rows': rows, 'graph': {'nodes': list(nodes.values()), 'edges': sorted(edges)}, 'summary': {
        'count': len(rows), 'direct': sum(r['impact'] == '직접' for r in rows),
        'indirect': sum(r['impact'] == '간접' for r in rows),
        'dev': dev, 'test': test, 'buffer': buf, 'total': total,
        'people': people, 'days': days, 'start': s.isoformat(), 'end': end.isoformat()}}

def write_xlsx(res):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook(); H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='305496')
    thin = Side(style='thin', color='BFBFBF'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    def sheet(ws, header, data, widths):
        ws.append(header)
        for c in ws[1]: c.font = H; c.fill = HF; c.alignment = Alignment(horizontal='center'); c.border = B
        for row in data: ws.append(row)
        for row in ws.iter_rows(min_row=2):
            for c in row: c.border = B
        for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = 'A2'
    sm = res['summary']; ws = wb.active; ws.title = '요약'
    sheet(ws, ['항목', '값'], [['변경 대상', res['query']], ['분석 일자', dt.date.today().isoformat()],
        ['영향 항목 수', sm['count']], ['직접 영향', sm['direct']], ['간접 영향', sm['indirect']],
        ['개발 공수 (MD)', sm['dev']], [f'테스트 공수 (MD, {int(TEST_RATIO*100)}%)', sm['test']],
        [f'버퍼 (MD, {int(BUFFER_RATIO*100)}%)', sm['buffer']], ['총 공수 (MD)', sm['total']],
        ['투입 인원', sm['people']], ['예상 기간 (영업일)', sm['days']], ['시작일', sm['start']], ['종료일', sm['end']]], [28, 20])
    ws2 = wb.create_sheet('영향목록')
    sheet(ws2, ['No', '구분', '항목', '파일', '영향', '경유', '공수(MD)', '시작일', '종료일', '담당자', '조치내용', '비고'],
          [[i + 1, r['kind'], r['name'], r['path'], r['impact'], r['via'], r['md'], r['start'], r['end'], r.get('owner', ''), r.get('action', ''), r.get('note', '')]
           for i, r in enumerate(res['rows'])], [5, 11, 40, 55, 7, 22, 9, 11, 11, 10, 30, 20])
    ws2.auto_filter.ref = ws2.dimensions
    g = res.get('graph') or {'nodes': [], 'edges': []}; lab = {n['id']: n for n in g['nodes']}
    ws4 = wb.create_sheet('연결관계')
    sheet(ws4, ['From 구분', 'From', 'To 구분', 'To', '관계'],
          [[lab[a]['kind'], lab[a]['label'], lab[b]['kind'], lab[b]['label'], 'include' if r else '참조' if lab[b]['kind'] == 'Table' else '호출'] for a, b, r in g['edges'] if a in lab and b in lab], [11, 40, 11, 40, 9])
    ws4.auto_filter.ref = ws4.dimensions
    ws3 = wb.create_sheet('공수기준')
    sheet(ws3, ['구분', '기본 MD/항목', '비고'], [[k, v, ''] for k, v in BASE_MD.items()] +
          [['파일크기 가중', '1.0 / 1.3 / 1.6', '<300 / <1000 / 1000+ 라인'], ['간접영향', 'x0.5', ''],
           ['테스트', f'x{TEST_RATIO}', '개발공수 대비'], ['버퍼', f'x{BUFFER_RATIO}', '개발+테스트 대비']], [16, 16, 30])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

HTML = r'''<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>영향도 분석</title>
<meta name="viewport" content="width=device-width,initial-scale=1">
<style>
:root{--bg:#eef1f5;--panel:#fff;--ink:#1c2430;--muted:#66717f;--line:#d5dbe3;--acc:#1f4e8c;--acc2:#e8f0fb;--direct:#b23a2a;--indirect:#8a6d1f}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font:14px/1.5 Pretendard,"Malgun Gothic",system-ui,sans-serif}
header{background:var(--acc);color:#fff;padding:14px 24px;display:flex;align-items:baseline;gap:14px}
header h1{margin:0;font-size:18px;font-weight:600}header span{opacity:.75;font-size:13px}
main{display:grid;grid-template-columns:300px 1fr;gap:16px;padding:16px 24px;min-height:calc(100vh - 52px)}
@media(max-width:820px){main{grid-template-columns:1fr}}
.panel{background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:16px}
label{display:block;font-size:12px;color:var(--muted);margin:10px 0 4px}label:first-child{margin-top:0}
input{width:100%;padding:8px 10px;border:1px solid var(--line);border-radius:4px;font:inherit}input:focus{outline:2px solid var(--acc);outline-offset:1px;border-color:var(--acc)}
.row{display:flex;gap:6px}.row input{flex:1}
button{font:inherit;padding:8px 14px;border-radius:4px;border:1px solid var(--acc);background:var(--acc);color:#fff;cursor:pointer}
button.ghost{background:#fff;color:var(--acc)}button:disabled{opacity:.5;cursor:default}
.actions{display:flex;gap:8px;margin-top:16px}.hint{font-size:12px;color:var(--muted);margin-top:14px;line-height:1.6}
.msg{margin-top:12px;padding:8px 10px;border-radius:4px;font-size:13px;display:none}.msg.err{display:block;background:#fbeaea;color:#8a2a1e}.msg.ok{display:block;background:var(--acc2);color:var(--acc)}
.stats{display:grid;grid-template-columns:repeat(auto-fit,minmax(120px,1fr));gap:10px;margin-bottom:14px}
.stat{padding:10px 12px;border:1px solid var(--line);border-radius:4px}.stat b{display:block;font-size:22px;font-weight:600;line-height:1.2}.stat span{font-size:12px;color:var(--muted)}
.stat.total{background:var(--acc2);border-color:var(--acc)}
.toolbar{display:flex;gap:8px;align-items:center;margin-bottom:10px;flex-wrap:wrap}.toolbar select,.toolbar input{padding:6px 8px;border:1px solid var(--line);border-radius:4px;font:inherit;width:auto}
.tw{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:13px}th,td{border-bottom:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top;white-space:nowrap}
th{position:sticky;top:0;background:#f6f8fb;font-weight:600;color:var(--muted);font-size:12px}
td.path{color:var(--muted);font-size:12px;max-width:380px;overflow:hidden;text-overflow:ellipsis}
.tag{display:inline-block;padding:1px 7px;border-radius:3px;font-size:12px}.tag.direct{background:#fbe9e6;color:var(--direct)}.tag.indirect{background:#fbf3df;color:var(--indirect)}
td input{padding:4px 6px;font-size:12px;min-width:90px}
.empty{padding:60px 0;text-align:center;color:var(--muted)}
.hist{margin-top:18px;border-top:1px solid var(--line);padding-top:12px}.hist h2{font-size:12px;color:var(--muted);font-weight:600;margin:0 0 6px}
.hist ul{list-style:none;margin:0;padding:0;max-height:260px;overflow:auto}.hist li{display:flex;align-items:center;gap:6px;padding:6px 8px;border-radius:4px;cursor:pointer;font-size:12px}
.hist li:hover{background:var(--acc2)}.hist li b{font-weight:600;color:var(--ink)}.hist li .r{color:var(--muted);overflow:hidden;text-overflow:ellipsis;white-space:nowrap;flex:1;min-width:0}
.hist li .m{color:var(--muted);white-space:nowrap}.hist li .x{border:0;background:none;color:var(--muted);padding:0 4px;cursor:pointer;font-size:14px;line-height:1}.hist li .x:hover{color:var(--direct)}
.hist .none{font-size:12px;color:var(--muted);padding:4px 8px}
.prog{border:1px solid var(--line);border-radius:4px;margin-bottom:14px;background:#fbfcfe}.prog .ph{display:flex;align-items:center;gap:8px;padding:8px 12px;font-size:13px;cursor:pointer;user-select:none}
.prog .ph b{font-weight:600}.prog .ph span{color:var(--muted);font-size:12px;margin-left:auto}
.prog.running .ph::before{content:'';width:10px;height:10px;border:2px solid var(--acc);border-top-color:transparent;border-radius:50%;animation:sp .8s linear infinite;flex:none}@keyframes sp{to{transform:rotate(360deg)}}
.prog pre{margin:0;padding:8px 12px;border-top:1px solid var(--line);font:12px/1.65 Consolas,"D2Coding",monospace;max-height:280px;overflow:auto;white-space:pre-wrap;color:var(--muted)}
.prog pre b{color:var(--ink);font-weight:600}.prog pre i{color:var(--direct);font-style:normal}.prog.fold pre{display:none}
.tabs{display:inline-flex;border:1px solid var(--line);border-radius:4px;overflow:hidden}.tabs button{border:0;border-radius:0;background:#fff;color:var(--muted);padding:6px 12px}.tabs button.on{background:var(--acc);color:#fff}
.gw{overflow:auto;border:1px solid var(--line);border-radius:4px;background:#fff;position:relative}.gw svg{display:block;font:12px Pretendard,"Malgun Gothic",system-ui,sans-serif}
.gw .col{fill:var(--muted);font-size:11px;font-weight:600;letter-spacing:.3px}.gw .colline{stroke:var(--line);stroke-dasharray:2 4}
.gw .n rect{fill:#fff;stroke:var(--line);stroke-width:1.2;rx:4}.gw .n text{fill:var(--ink)}.gw .n{cursor:pointer}
.gw .n.direct rect{stroke:var(--direct);fill:#fdf3f1}.gw .n.indirect rect{stroke:#d9b45a;fill:#fdf8ea}
.gw .e{fill:none;stroke:#9aa7b8;stroke-width:1.4;marker-end:url(#ar)}.gw .e.include{stroke-dasharray:4 3}
.gw.focus .n:not(.hl){opacity:.18}.gw.focus .e:not(.hl){opacity:.08}.gw .e.hl{stroke:var(--acc);stroke-width:2.2;marker-end:url(#arh)}.gw .n.hl rect{stroke-width:2;filter:drop-shadow(0 1px 2px rgba(0,0,0,.15))}.gw .n.hl.pin rect{stroke:var(--acc)}
.ghint{font-size:12px;color:var(--muted);margin:8px 0 0;display:flex;gap:14px;flex-wrap:wrap}.ghint i{display:inline-block;width:10px;height:10px;border-radius:2px;vertical-align:-1px;margin-right:4px;border:1.5px solid}
</style></head><body>
<header><h1>영향도 분석</h1><span>Spring MVC · JSP/JS → Controller → Service → Mapper → 테이블</span></header>
<main>
<section class="panel">
 <label for="root">소스 루트 폴더</label>
 <div class="row"><input id="root" list="roots" placeholder="D:\workspace\dWorks\src"><button class="ghost" id="pick" title="폴더 선택">…</button></div>
 <label for="q">검색어</label><input id="q" list="queries" placeholder="TB_USER / selectUserList / /user/list / fn_search">
 <label for="start">작업 시작일</label><input id="start" type="date">
 <label for="people">투입 인원</label><input id="people" type="number" step="0.5" min="0.5" value="1">
 <div class="actions"><button id="run">분석 실행</button><button class="ghost" id="export" disabled>엑셀로 저장</button></div>
 <div class="msg" id="msg"></div>
 <datalist id="roots"></datalist><datalist id="queries"></datalist>
 <div class="hist"><h2>최근 분석</h2><ul id="hist"></ul></div>
 <div class="hint">검색어는 테이블명, mapper id, 컨트롤러 URL 일부, JS 함수명(fn_xxx) 모두 가능합니다. 대소문자 구분 없이 부분 일치로 찾습니다.<br><br>표에서 담당자·조치내용·비고를 입력하면 엑셀에 같이 저장됩니다.</div>
</section>
<section class="panel">
 <div id="prog" class="prog" hidden><div class="ph"><b id="ptitle"></b><span id="ptime"></span></div><pre id="plog"></pre></div>
 <div id="stats" class="stats"></div>
 <div class="toolbar">
  <div class="tabs"><button class="on" data-v="table">표</button><button data-v="graph">관계도</button></div>
  <select id="fkind"><option value="">구분: 전체</option><option>Table</option><option>Mapper</option><option>DAO</option><option>Service</option><option>Controller</option><option>JSP</option><option>JS</option></select>
  <select id="fimpact"><option value="">영향: 전체</option><option>직접</option><option>간접</option></select>
  <input id="ftext" placeholder="항목·파일 필터">
  <span id="fcount" style="color:var(--muted);font-size:12px"></span>
 </div>
 <div class="tw"><table><thead><tr><th>No</th><th>구분</th><th>항목</th><th>파일</th><th>영향</th><th>경유</th><th>MD</th><th>시작</th><th>종료</th><th>담당자</th><th>조치내용</th><th>비고</th></tr></thead>
 <tbody id="tb"><tr><td colspan="12" class="empty">소스 루트와 검색어를 입력하고 분석을 실행하세요.</td></tr></tbody></table></div>
 <div id="gwrap" hidden><div class="gw" id="gw"></div>
  <div class="ghint"><span><i style="border-color:var(--direct);background:#fdf3f1"></i>직접</span><span><i style="border-color:#d9b45a;background:#fdf8ea"></i>간접</span><span>실선: 호출 · 점선: script include</span><span>노드에 마우스를 올리면 연결된 경로만 강조, 클릭하면 고정</span></div></div>
</section>
</main>
<script>
const $=s=>document.querySelector(s);let res=null;
$('#start').value=new Date().toISOString().slice(0,10);
const esc=s=>String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
function msg(t,cls){const m=$('#msg');m.textContent=t;m.className='msg '+(cls||'');}
function stats(){const s=res.summary;$('#stats').innerHTML=[
 ['영향 항목',s.count+'건',`직접 ${s.direct} · 간접 ${s.indirect}`],['개발 공수',s.dev+' MD',''],['테스트 · 버퍼',(s.test+s.buffer).toFixed(2)+' MD',`${s.test} + ${s.buffer}`],
 ['총 공수',s.total+' MD',`${s.people}명 투입`,'total'],['예상 기간',s.days+' 영업일',`${s.start} ~ ${s.end}`]
].map(([l,v,sub,c])=>`<div class="stat ${c||''}"><span>${l}</span><b>${v}</b><span>${sub}</span></div>`).join('');}
const COLS=['JSP','JS','Controller','Service','DAO','Mapper','Table'];
function graph(){const g=res.graph,gw=$('#gw');gw.className='gw';
 if(!g.nodes.length){gw.innerHTML='<div class="empty">표시할 관계가 없습니다.</div>';return;}
 const N={};g.nodes.forEach(n=>{N[n.id]=n;n.c=Math.max(0,COLS.indexOf(n.kind));n.out=[];n.in=[];});
 const E=g.edges.filter(([a,b])=>N[a]&&N[b]);E.forEach(([a,b])=>{N[a].out.push(b);N[b].in.push(a);});
 const cols=COLS.map(()=>[]);g.nodes.forEach(n=>cols[n.c].push(n));
 const idx={};const reidx=()=>cols.forEach(c=>c.forEach((n,i)=>idx[n.id]=i));reidx();
 for(let p=0;p<4;p++){const dir=p%2?[...cols].reverse():cols;dir.forEach(c=>{c.forEach(n=>{const nb=[...n.in,...n.out].map(x=>idx[x]);n.k=nb.length?nb.reduce((a,b)=>a+b,0)/nb.length:idx[n.id];});c.sort((a,b)=>a.k-b.k||a.label.localeCompare(b.label));reidx();});}
 const CW=270,NW=190,NH=30,GAP=12,TOP=40,LEFT=16,rows=Math.max(...cols.map(c=>c.length)),H=TOP+rows*(NH+GAP)+16,W=LEFT*2+COLS.length*CW-(CW-NW);
 cols.forEach((c,ci)=>{const off=(rows-c.length)*(NH+GAP)/2;c.forEach((n,i)=>{n.x=LEFT+ci*CW;n.y=TOP+off+i*(NH+GAP);});});
 const cut=(t,n)=>t.length>n?t.slice(0,n-1)+'…':t;
 // 한 노드에 선이 여러 개 붙으면 접점을 세로로 조금씩 벌려서 겹치지 않게
 const port=(n,list,key)=>{const arr=list.slice().sort((p,q)=>N[p[key]].y-N[q[key]].y),step=Math.min(7,18/Math.max(1,arr.length-1));const m={};arr.forEach((e,i)=>m[e[3]]=n.y+NH/2+(i-(arr.length-1)/2)*step);return m;};
 E.forEach((e,i)=>e[3]=i);const outP={},inP={};
 g.nodes.forEach(n=>{outP[n.id]=port(n,E.filter(e=>e[0]===n.id),1);inP[n.id]=port(n,E.filter(e=>e[1]===n.id),0);});
 let svg=`<svg width="${W}" height="${H}" viewBox="0 0 ${W} ${H}"><defs>
  <marker id="ar" viewBox="0 0 8 8" refX="7" refY="4" markerWidth="7" markerHeight="7" orient="auto-start-reverse"><path d="M0 0.5 L7 4 L0 7.5 z" fill="#9aa7b8"/></marker>
  <marker id="arh" viewBox="0 0 8 8" refX="7" refY="4" markerWidth="7" markerHeight="7" orient="auto-start-reverse"><path d="M0 0.5 L7 4 L0 7.5 z" fill="#1f4e8c"/></marker></defs>`;
 COLS.forEach((k,ci)=>{svg+=`<text class="col" x="${LEFT+ci*CW}" y="22">${k}${cols[ci].length?` (${cols[ci].length})`:''}</text>`;if(ci)svg+=`<line class="colline" x1="${LEFT+ci*CW-(CW-NW)/2}" y1="30" x2="${LEFT+ci*CW-(CW-NW)/2}" y2="${H-8}"/>`;});
 E.forEach(e=>{const [a,b,r,i]=e,A=N[a],B=N[b],y1=outP[a][i],y2=inP[b][i];let d;
  if(B.c>A.c){const x1=A.x+NW,x2=B.x,k=Math.max(40,(x2-x1)*0.5);d=`M${x1} ${y1} C${x1+k} ${y1}, ${x2-k} ${y2}, ${x2} ${y2}`;}          // 오른쪽으로: 오른쪽 변 → 왼쪽 변
  else if(B.c<A.c){const x1=A.x,x2=B.x+NW,k=Math.max(40,(x1-x2)*0.5);d=`M${x1} ${y1} C${x1-k} ${y1}, ${x2+k} ${y2}, ${x2} ${y2}`;}     // 왼쪽으로(역방향): 왼쪽 변 → 오른쪽 변
  else{const x=A.x,k=Math.min(50,(CW-NW)*0.7);d=`M${x} ${y1} C${x-k} ${y1}, ${x-k} ${y2}, ${x} ${y2}`;}                                 // 같은 열(JS→JS 호출): 왼쪽으로 볼록한 곡선
  svg+=`<path class="e ${r}" data-a="${esc(a)}" data-b="${esc(b)}" d="${d}"/>`;});
 g.nodes.forEach(n=>{svg+=`<g class="n ${n.impact==='직접'?'direct':'indirect'}" data-id="${esc(n.id)}" transform="translate(${n.x},${n.y})"><title>${esc(n.label)}${n.title?'\n'+esc(n.title):''}</title><rect width="${NW}" height="${NH}"/><text x="9" y="19">${esc(cut(n.label,27))}</text></g>`;});
 gw.innerHTML=svg+'</svg>';
 const reach=(id,dir)=>{const seen=new Set([id]),st=[id];while(st.length){const x=st.pop();N[x][dir].forEach(y=>{if(!seen.has(y)){seen.add(y);st.push(y);}});}return seen;};
 let pin=null;
 const focus=id=>{gw.classList.toggle('focus',!!id);gw.querySelectorAll('.hl').forEach(e=>e.classList.remove('hl','pin'));if(!id)return;
  const set=new Set([...reach(id,'in'),...reach(id,'out')]);
  gw.querySelectorAll('.n').forEach(e=>{if(set.has(e.dataset.id))e.classList.add('hl');});gw.querySelector(`.n[data-id="${CSS.escape(id)}"]`).classList.add('pin');
  gw.querySelectorAll('.e').forEach(e=>{if(set.has(e.dataset.a)&&set.has(e.dataset.b))e.classList.add('hl');});};
 gw.onmouseover=e=>{const n=e.target.closest('.n');if(n&&!pin)focus(n.dataset.id);};
 gw.onmouseout=e=>{if(!pin&&e.target.closest('.n'))focus(null);};
 gw.onclick=e=>{const n=e.target.closest('.n');if(!n){pin=null;focus(null);return;}pin=pin===n.dataset.id?null:n.dataset.id;focus(pin||null);};}
document.querySelectorAll('.tabs button').forEach(b=>b.onclick=()=>{document.querySelectorAll('.tabs button').forEach(x=>x.classList.toggle('on',x===b));const g=b.dataset.v==='graph';
 $('.tw').hidden=g;$('#gwrap').hidden=!g;['#fkind','#fimpact','#ftext','#fcount'].forEach(s=>$(s).style.visibility=g?'hidden':'');if(g&&res)graph();});
function render(){const k=$('#fkind').value,i=$('#fimpact').value,t=$('#ftext').value.toLowerCase();
 const rows=res.rows.map((r,idx)=>({r,idx})).filter(({r})=>(!k||r.kind===k)&&(!i||r.impact===i)&&(!t||(r.name+r.path+r.via).toLowerCase().includes(t)));
 $('#fcount').textContent=`${rows.length} / ${res.rows.length}`;
 if(!rows.length){$('#tb').innerHTML='<tr><td colspan="12" class="empty">일치하는 항목이 없습니다.</td></tr>';return;}
 $('#tb').innerHTML=rows.map(({r,idx})=>`<tr><td>${idx+1}</td><td>${r.kind}</td><td>${esc(r.name)}</td><td class="path" title="${esc(r.path)}">${esc(r.path)}</td>
 <td><span class="tag ${r.impact==='직접'?'direct':'indirect'}">${r.impact}</span></td><td>${esc(r.via)}</td><td>${r.md}</td><td>${r.start}</td><td>${r.end}</td>
 <td><input data-i="${idx}" data-f="owner" value="${esc(r.owner)}"></td><td><input data-i="${idx}" data-f="action" value="${esc(r.action)}"></td><td><input data-i="${idx}" data-f="note" value="${esc(r.note)}"></td></tr>`).join('');}
$('#tb').addEventListener('input',e=>{const d=e.target.dataset;if(d.i!==undefined)res.rows[d.i][d.f]=e.target.value;});
['#fkind','#fimpact','#ftext'].forEach(s=>$(s).addEventListener('input',()=>res&&render()));
const sleep=ms=>new Promise(z=>setTimeout(z,ms));
function prog(log,ms,state){const p=$('#prog');p.hidden=false;p.className='prog '+state;
 $('#ptitle').textContent={running:'분석 중…',ok:'분석 완료',err:'분석 실패'}[state];$('#ptime').textContent=(ms/1000).toFixed(1)+'초';
 $('#plog').innerHTML=log.map((l,i)=>{const h=esc(l);return l.includes('✖')?`<i>${h}</i>`:i===log.length-1&&state==='running'?`<b>${h}</b>`:h;}).join('\n');$('#plog').scrollTop=1e9;}
$('#prog').querySelector('.ph').onclick=()=>$('#prog').classList.toggle('fold');
$('#run').onclick=async()=>{const body={root:$('#root').value.trim(),query:$('#q').value.trim(),start:$('#start').value,people:parseFloat($('#people').value)||1};
 if(!body.root||!body.query){msg('소스 루트와 검색어를 입력하세요.','err');return;}
 $('#run').disabled=true;msg('분석 중…','ok');const t0=Date.now();let log=[],ok=false;
 try{const r=await fetch('/run',{method:'POST',body:JSON.stringify(body)});if(!r.ok)throw new Error((await r.json()).error);
  let j;while(true){j=await(await fetch('/progress')).json();log=j.log;prog(log,Date.now()-t0,'running');if(j.done)break;await sleep(250);}
  if(j.error)throw new Error(j.error);
  res=j.result;res.rows.forEach(x=>{x.owner='';x.action='';x.note='';});hist=j.result.history||hist;renderHist();stats();render();if(!$('#gwrap').hidden)graph();$('#export').disabled=false;ok=true;
  msg(res.rows.length?`${res.rows.length}건 찾았습니다.`:'0건입니다. 검색어와 소스 루트를 확인하세요.','ok');
 }catch(e){msg('분석 실패: '+e.message,'err');}finally{$('#run').disabled=false;prog(log,Date.now()-t0,ok?'ok':'err');if(ok)$('#prog').classList.add('fold');}};
$('#export').onclick=async()=>{const r=await fetch('/export',{method:'POST',body:JSON.stringify(res)});if(!r.ok){msg('엑셀 저장 실패: '+(await r.json()).error,'err');return;}
 const b=await r.blob();const a=document.createElement('a');a.href=URL.createObjectURL(b);a.download=`영향도분석서_${res.query}_${res.summary.start}.xlsx`;a.click();msg('엑셀을 저장했습니다.','ok');};
$('#pick').onclick=async()=>{const r=await fetch('/pick');const j=await r.json();if(j.path)$('#root').value=j.path;else if(j.error)msg(j.error,'err');};
let hist=[];
function renderHist(){$('#roots').innerHTML=[...new Set(hist.map(h=>h.root))].map(r=>`<option value="${esc(r)}">`).join('');
 $('#queries').innerHTML=[...new Set(hist.map(h=>h.query))].map(q=>`<option value="${esc(q)}">`).join('');
 $('#hist').innerHTML=hist.length?hist.map((h,i)=>`<li data-i="${i}" title="${esc(h.root)}"><b>${esc(h.query)}</b><span class="r">${esc(h.root)}</span><span class="m">${h.count}건 · ${h.total}MD · ${h.at}</span><button class="x" data-del="${i}" title="삭제">×</button></li>`).join(''):'<li class="none">아직 없습니다. 분석을 실행하면 여기에 쌓입니다.</li>';}
$('#hist').addEventListener('click',async e=>{const d=e.target.dataset;
 if(d.del!==undefined){const h=hist[d.del];const r=await fetch('/history/delete',{method:'POST',body:JSON.stringify({root:h.root,query:h.query})});hist=await r.json();renderHist();return;}
 const li=e.target.closest('li');if(!li||li.dataset.i===undefined)return;const h=hist[li.dataset.i];
 $('#root').value=h.root;$('#q').value=h.query;$('#people').value=h.people;$('#q').focus();});
fetch('/history').then(r=>r.json()).then(j=>{hist=j;renderHist();if(hist[0]){$('#root').value=hist[0].root;$('#people').value=hist[0].people;}});
$('#q').addEventListener('keydown',e=>{if(e.key==='Enter')$('#run').click();});
</script></body></html>'''


# ---------- 분석 작업 (백그라운드 스레드) ----------
JOB = {'log': [], 'done': True, 'result': None, 'error': None}   # log: [(replace, text)]
JOB_LOCK = threading.Lock(); _cr = [False]
try: sys.stdout.reconfigure(errors='replace')   # cp949 콘솔에서 ✔ 같은 기호 때문에 죽지 않게
except Exception: pass

def job_log(s, replace=False):
    if replace: print('\r' + s[:110].ljust(110), end='', flush=True); _cr[0] = True
    else:
        if _cr[0]: print(); _cr[0] = False
        print(s, flush=True)
    with JOB_LOCK:
        if replace and JOB['log'] and JOB['log'][-1][0]: JOB['log'][-1] = (True, s)
        else: JOB['log'].append((replace, s))

def job_start(b):
    with JOB_LOCK:
        if not JOB['done']: raise ValueError('이미 분석이 진행 중입니다. 끝날 때까지 기다려 주세요.')
        JOB.update(log=[], done=False, result=None, error=None)
    def work():
        try:
            people = float(b.get('people') or 1)
            res = analyze(b['root'], b['query'], b['start'], people, log=job_log)
            res['history'] = save_hist({'root': b['root'], 'query': b['query'], 'people': people,
                                        'count': res['summary']['count'], 'total': res['summary']['total'], 'at': dt.datetime.now().strftime('%m-%d %H:%M')})
            with JOB_LOCK: JOB['result'] = res
        except Exception as e:
            job_log(f'   ✖ 실패: {e}')
            with JOB_LOCK: JOB['error'] = str(e)
        finally:
            with JOB_LOCK: JOB['done'] = True
    threading.Thread(target=work, daemon=True).start()

def job_state():
    with JOB_LOCK:
        return {'log': [x[1] for x in JOB['log']], 'done': JOB['done'], 'error': JOB['error'],
                'result': JOB['result'] if JOB['done'] else None}

class H(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def _send(self, code, body, ctype='application/json; charset=utf-8', extra=None):
        self.send_response(code); self.send_header('Content-Type', ctype)
        for k, v in (extra or {}).items(): self.send_header(k, v)
        self.end_headers(); self.wfile.write(body)
    def _json(self):
        return json.loads(self.rfile.read(int(self.headers.get('Content-Length', 0))) or b'{}')
    def do_GET(self):
        p = urlparse(self.path).path
        if p == '/': return self._send(200, HTML.encode('utf-8'), 'text/html; charset=utf-8')
        if p == '/history': return self._send(200, json.dumps(load_hist(), ensure_ascii=False).encode('utf-8'))
        if p == '/progress': return self._send(200, json.dumps(job_state(), ensure_ascii=False).encode('utf-8'))
        if p == '/pick':
            try:
                import tkinter as tk; from tkinter import filedialog
                r = tk.Tk(); r.withdraw(); r.attributes('-topmost', True)
                d = filedialog.askdirectory(title='소스 루트 폴더 선택'); r.destroy()
                return self._send(200, json.dumps({'path': d.replace('/', '\\') if d else ''}).encode())
            except Exception as e:
                return self._send(200, json.dumps({'error': f'폴더 선택창을 열 수 없습니다 ({e}). 경로를 직접 입력하세요.'}).encode())
        self._send(404, b'{}')
    def do_POST(self):
        p = urlparse(self.path).path
        try:
            b = self._json()
            if p == '/run':
                job_start(b)
                return self._send(200, b'{"started":true}')
            if p == '/history/delete':
                h = [x for x in load_hist() if not (x['root'] == b['root'] and x['query'] == b['query'])]
                HIST_FILE.write_text(json.dumps(h, ensure_ascii=False, indent=1), encoding='utf-8')
                return self._send(200, json.dumps(h, ensure_ascii=False).encode('utf-8'))
            if p == '/export':
                return self._send(200, write_xlsx(b), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                                  {'Content-Disposition': 'attachment; filename="impact.xlsx"'})
            self._send(404, b'{}')
        except ModuleNotFoundError:
            self._send(500, json.dumps({'error': 'openpyxl이 없습니다. python -m pip install openpyxl'}, ensure_ascii=False).encode('utf-8'))
        except Exception as e:
            self._send(400, json.dumps({'error': str(e)}, ensure_ascii=False).encode('utf-8'))

if __name__ == '__main__':
    srv = HTTPServer(('127.0.0.1', PORT), H)
    url = f'http://127.0.0.1:{PORT}'
    print(f'영향도 분석 UI: {url}  (종료: Ctrl+C)')
    threading.Timer(0.5, lambda: webbrowser.open(url)).start()
    try: srv.serve_forever()
    except KeyboardInterrupt: pass
